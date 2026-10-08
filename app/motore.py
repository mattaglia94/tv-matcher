"""Flusso completo: legge input Excel -> Amazon -> retailer -> scrive output Excel."""
import datetime as dt
import os
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config_default import RETAILER
from nucleo import ESITI

# ------------------------------------------------------------ lettura input
def leggi_input(percorso):
    wb = load_workbook(percorso, data_only=True)
    ws = wb["ASIN"] if "ASIN" in wb.sheetnames else wb.worksheets[0]
    righe = []
    intest = [str(c.value or "").strip().lower() for c in ws[1]]

    def col(*nomi):
        for n in nomi:
            if n in intest:
                return intest.index(n)
        return None
    c_asin = col("asin")
    c_nome = col("amazon product name", "nome prodotto amazon", "item name")
    c_mod = col("amazon model", "modello amazon")
    c_paese = col("paese", "country")
    start = 2
    if c_asin is None:  # foglio senza intestazione: ASIN in colonna A
        c_asin, start = 0, 1
    for r in ws.iter_rows(min_row=start, values_only=True):
        a = str(r[c_asin] or "").strip().upper()
        if not re.fullmatch(r"B0[A-Z0-9]{8}", a):
            continue
        righe.append({"asin": a,
                      "nome": str(r[c_nome] or "").strip() if c_nome is not None else "",
                      "modello": str(r[c_mod] or "").strip() if c_mod is not None else "",
                      "paese": str(r[c_paese] or "").strip().upper() if c_paese is not None else ""})
    retailer = []
    if "Retailer" in wb.sheetnames:
        for r in wb["Retailer"].iter_rows(min_row=2, values_only=True):
            if r and r[0] and r[1] and r[2]:
                retailer.append((str(r[0]).strip().upper(), str(r[1]).strip(), str(r[2]).strip(),
                                 str(r[3] or "SI").strip().upper()))
    impost = {"anno": 2026, "anno_ignoto": True, "verifica": True}
    if "Impostazioni" in wb.sheetnames:
        for r in wb["Impostazioni"].iter_rows(min_row=2, values_only=True):
            if not r or not r[0]:
                continue
            k, v = str(r[0]).lower(), r[1]
            if "anno target" in k:
                impost["anno"] = int(v)
            elif "anno non" in k:
                impost["anno_ignoto"] = str(v).upper().startswith("S")
            elif "verifica" in k:
                impost["verifica"] = str(v).upper().startswith("S")
    return righe, retailer or list(RETAILER), impost


# ------------------------------------------------------------ esecuzione
def esegui(percorso_input, paese, cartella_out, visibile=False, log=print, stop=lambda: False):
    from browser import Browser
    from nucleo import elabora_asin
    righe, retailer, imp = leggi_input(percorso_input)
    paese = paese.upper()
    ret = [{"nome": n, "url": u} for (p, n, u, att) in retailer if p == paese and att.startswith("S")]
    righe = [r for r in righe if not r["paese"] or r["paese"] == paese]
    log(f"{len(righe)} ASIN, {len(ret)} retailer per {paese}: {', '.join(r['nome'] for r in ret)}")
    if not righe or not ret:
        raise RuntimeError("Nessun ASIN o nessun retailer attivo per il paese scelto.")
    parametri = {"anno": {"attivo": True, "valore": imp["anno"], "accetta_ignoto": imp["anno_ignoto"]},
                 "verifica_url": imp["verifica"]}
    br = Browser(visibile=visibile, log=log)
    risultati = []
    try:
        for i, r in enumerate(righe, 1):
            if stop():
                break
            log(f"[{i}/{len(righe)}] {r['asin']}")
            r["paese"] = paese
            risultati.append(elabora_asin(br, r, ret, parametri, log=log, stop=stop))
    finally:
        br.chiudi()
    return scrivi_output(risultati, [r["nome"] for r in ret], paese, cartella_out)


# ------------------------------------------------------------ output
def scrivi_output(risultati, nomi_ret, paese, cartella_out):
    wb = Workbook()
    ws = wb.active
    ws.title = f"Matching {paese}"
    blu = PatternFill("solid", fgColor="1F4E78")
    azz = PatternFill("solid", fgColor="DDEBF7")
    bianco = Font(bold=True, color="FFFFFF")
    ws.append(["ASIN", "Amazon product name", "Amazon Model"] +
              [x for n in nomi_ret for x in (n, "")])
    ws.append(["", "", ""] + [x for _ in nomi_ret for x in ("Model Name", "Product URL")])
    for c in range(1, 4):
        ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)
    for i, _ in enumerate(nomi_ret):
        c = 4 + 2 * i
        ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=c + 1)
    for row in ws.iter_rows(min_row=1, max_row=2):
        for cell in row:
            cell.font = bianco if cell.row == 1 else Font(bold=True)
            cell.fill = blu if cell.row == 1 else azz
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in risultati:
        riga = [r["asin"], r["amazon"].get("titolo", ""), r["amazon"].get("modello", "")]
        for n in nomi_ret:
            m = (r.get("match") or {}).get(n) or []
            esito = (r.get("esito") or {}).get(n)
            if m:
                riga += [" / ".join(x["modello"] for x in m), " ".join(x["url"] for x in m)]
            elif esito and esito != "trovato":
                riga += [ESITI[esito], ""]
            else:
                riga += ["", ""]
        ws.append(riga)
        rr = ws.max_row
        for i, n in enumerate(nomi_ret):
            cell = ws.cell(rr, 5 + 2 * i)
            if cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
    ws.column_dimensions["A"].width = 13
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["C"].width = 18
    for i in range(len(nomi_ret)):
        ws.column_dimensions[get_column_letter(4 + 2 * i)].width = 18
        ws.column_dimensions[get_column_letter(5 + 2 * i)].width = 45
    ws.freeze_panes = "D3"

    wd = wb.create_sheet("Dettaglio")
    wd.append(["ASIN", "Fonte", "Modello trovato", "Famiglia", "Pollici", "Anno", "URL", "Esito", "Ricerca"])
    for c in wd[1]:
        c.font, c.fill = bianco, blu
    for r in risultati:
        a = r["amazon"]
        wd.append([r["asin"], "Amazon", a.get("modello"), a.get("famiglia"), a.get("pollici"), a.get("anno"),
                   a.get("url"), r.get("errore") or "riferimento", ""])
        for d in r["dettaglio"]:
            wd.append([r["asin"], d.get("retailer"), d.get("modello", ""), d.get("famiglia", ""), d.get("pollici", ""),
                       d.get("anno", ""), d.get("url", ""), d.get("esito", ""), d.get("ricerca", "")])
    wd.auto_filter.ref = wd.dimensions
    for col, w in zip("ABCDEFGHI", (13, 16, 18, 30, 8, 7, 60, 34, 26)):
        wd.column_dimensions[col].width = w

    os.makedirs(cartella_out, exist_ok=True)
    nome = f"Matching_TV_{paese}_{dt.datetime.now():%Y-%m-%d_%H%M}.xlsx"
    out = os.path.join(cartella_out, nome)
    wb.save(out)
    return out


def crea_modello_input(percorso):
    """Crea il file di input vuoto con i 3 fogli."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ASIN"
    ws.append(["ASIN", "Paese", "Amazon product name", "Amazon Model"])
    for c in ws[1]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["D"].width = 18
    wr = wb.create_sheet("Retailer")
    wr.append(["Paese", "Nome retailer", "URL ricerca ({q} = testo cercato)", "Attivo (SI/NO)"])
    for c in wr[1]:
        c.font = Font(bold=True)
    for r in RETAILER:
        wr.append(list(r))
    wr.column_dimensions["B"].width = 20
    wr.column_dimensions["C"].width = 75
    wi = wb.create_sheet("Impostazioni")
    wi.append(["Impostazione", "Valore", "Spiegazione"])
    wi.append(["Anno target", 2026, "Solo modelli di quest'anno"])
    wi.append(["Accetta anno non determinabile", "SI",
               "SI = accetta modelli il cui anno non si ricava dal codice/titolo (segnalati nel foglio Dettaglio)"])
    wi.append(["Verifica pagina prodotto", "SI", "SI = apre ogni URL scelto per controllare che sia attivo"])
    for c in wi[1]:
        c.font = Font(bold=True)
    wi.column_dimensions["A"].width = 32
    wi.column_dimensions["C"].width = 90
    wb.save(percorso)
